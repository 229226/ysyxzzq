#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
把 sim.cpp 写的 build/perf_table.csv 复制进 Excel 的 "NPC性能评估结果" 表。

表格是两行等长的 CSV：第一行是字段名，第二行是对应的值。
第一列就对应 A 列，往后依次对齐，所以本脚本只是按位置搬过去，不做解析也不做映射。

A(commit)、B(说明)、F(综合频率)、G(综合面积) 仿真产不出来，sim.cpp 会占空位，
搬过来就是空的，填不填由你自己决定。

用法: python3 fill_excel.py <perf_table.csv> [excel_file]

注意：脚本不写 A(commit) / B(说明)，选行靠「第一个 A 列为空的行」，
      所以连着跑两次而没填 A 列会覆盖同一行，填完记得补上 A 列。
"""

import csv
import sys
import os
from openpyxl import load_workbook


def col_letter(idx):
    """0 -> A, 25 -> Z, 26 -> AA"""
    s = ''
    idx += 1
    while idx:
        idx, r = divmod(idx - 1, 26)
        s = chr(ord('A') + r) + s
    return s


def to_num(s):
    """能当数字就当数字写，别让整列变成文本；空串原样返回"""
    for cast in (int, float):
        try:
            return cast(s)
        except ValueError:
            pass
    return s


def main():
    if len(sys.argv) < 2:
        print("用法: python3 fill_excel.py <perf_table.csv> [excel_file]")
        sys.exit(1)

    table_file = sys.argv[1]
    excel_file = sys.argv[2] if len(sys.argv) > 2 else 'NPC性能评估结果.xlsx'

    if not os.path.isfile(table_file):
        print(f"错误: 找不到 {table_file}，先跑一次仿真（make sim）生成它", file=sys.stderr)
        sys.exit(1)

    with open(table_file, 'r', encoding='utf-8') as f:
        rows = list(csv.reader(f))

    if len(rows) < 2 or len(rows[0]) != len(rows[1]):
        print(f"错误: {table_file} 不是两行等长的性能表格", file=sys.stderr)
        sys.exit(1)

    values = rows[1]

    wb = load_workbook(excel_file)
    ws = wb['NPC性能评估结果']

    # 第一个 A 列为空的行
    row = 4
    while ws.cell(row=row, column=1).value is not None:
        row += 1

    for i, v in enumerate(values):
        if v != '':
            ws[col_letter(i) + str(row)] = to_num(v)

    wb.save(excel_file)
    print(f"{len(values)} 列已写入第 {row} 行，保存至 {excel_file}")


if __name__ == '__main__':
    main()
